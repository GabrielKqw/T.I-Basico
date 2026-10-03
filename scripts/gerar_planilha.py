"""Gera planilhas/exercicios-excel.xlsx usada nos exercícios do curso."""
import random
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

random.seed(42)
SAIDA = Path(__file__).resolve().parent.parent / "planilhas" / "exercicios-excel.xlsx"
CAB_FILL = PatternFill("solid", fgColor="1F4E79")
CAB_FONT = Font(bold=True, color="FFFFFF")
FAZER_FILL = PatternFill("solid", fgColor="FFF2CC")


def aba(wb, nome, cabecalho, linhas, colunas_fazer=()):
    ws = wb.create_sheet(nome)
    ws.append(cabecalho)
    for linha in linhas:
        ws.append(linha)
    for c in ws[1]:
        c.fill, c.font = CAB_FILL, CAB_FONT
        c.alignment = Alignment(horizontal="center")
    for col in colunas_fazer:  # colunas que o aluno deve preencher
        for r in range(2, len(linhas) + 2):
            ws[f"{col}{r}"].fill = FAZER_FILL
    for i, titulo in enumerate(cabecalho, 1):
        ws.column_dimensions[get_column_letter(i)].width = max(14, len(str(titulo)) + 4)
    ws.freeze_panes = "A2"
    return ws


wb = Workbook()
ws = wb.active
ws.title = "Instruções"
for linha in [
    ["Curso de T.I. Básico - Planilha de Exercícios"],
    [],
    ["Os enunciados estão no Módulo 11 do curso (Exercícios e Gabarito)."],
    ["Células em AMARELO devem ser preenchidas por você com fórmulas."],
    [],
    ["Aba", "Exercícios"],
    ["Notas", "Parte D - média, SE, MÁXIMO, MÍNIMO, CONT.SE"],
    ["Vendas", "Partes E e H - SOMASE, SOMASES, CONT.SE, tabela dinâmica, gráficos"],
    ["Produtos", "Tabela de consulta para PROCV/PROCX"],
    ["Pedidos", "Parte E - PROCV/PROCX e totais"],
    ["Limpeza", "Parte F - ARRUMAR, PRI.MAIÚSCULA, ESQUERDA, TEXTO"],
    ["Datas", "Parte G - DATADIF, HOJE, soma de dias"],
]:
    ws.append(linha)
ws["A1"].font = Font(bold=True, size=16, color="1F4E79")
ws["A6"].font = ws["B6"].font = Font(bold=True)
ws["A4"].fill = FAZER_FILL
ws.column_dimensions["A"].width = 20
ws.column_dimensions["B"].width = 70

# Notas
alunos = ["Ana Souza", "Bruno Lima", "Carla Dias", "Diego Alves", "Elisa Rocha",
          "Fábio Nunes", "Gabriela Melo", "Heitor Costa", "Isabela Reis", "João Pedro",
          "Karina Lopes", "Lucas Martins", "Mariana Silva", "Nicolas Ramos", "Olívia Torres"]
notas = [[a] + [round(random.uniform(2, 10), 1) for _ in range(4)] + [None, None] for a in alunos]
ws = aba(wb, "Notas", ["Aluno", "Nota 1", "Nota 2", "Nota 3", "Nota 4", "Média", "Situação"],
         notas, colunas_fazer="FG")
for r, rotulo in enumerate(["Maior média", "Menor média", "Média geral", "Aprovados"], 2):
    ws[f"I{r}"] = rotulo
    ws[f"I{r}"].font = Font(bold=True)
    ws[f"J{r}"].fill = FAZER_FILL
ws.column_dimensions["I"].width = 16

# Produtos
produtos = [
    ("P001", "Notebook", "Informática", 3500.00),
    ("P002", "Monitor", "Informática", 1200.00),
    ("P003", "Mouse", "Acessórios", 80.00),
    ("P004", "Teclado", "Acessórios", 150.00),
    ("P005", "Impressora", "Informática", 950.00),
    ("P006", "Headset", "Acessórios", 220.00),
    ("P007", "Webcam", "Acessórios", 310.00),
    ("P008", "Cadeira Gamer", "Móveis", 1450.00),
]
ws = aba(wb, "Produtos", ["Código", "Produto", "Categoria", "Preço"], produtos)
for r in range(2, 10):
    ws[f"D{r}"].number_format = '"R$" #,##0.00'

# Vendas
vendedores = {"Ana": "Sul", "Bruno": "Norte", "Carla": "Sudeste", "Diego": "Nordeste", "Elisa": "Centro-Oeste"}
regioes = list(vendedores.values()) + ["Sul", "Sudeste"]
vendas = []
inicio = date(2026, 7, 1)
for i in range(40):
    vend = random.choice(list(vendedores))
    reg = vendedores[vend] if random.random() < 0.7 else random.choice(regioes)
    _, prod, _, preco = random.choice(produtos)
    qtd = random.randint(1, 5)
    vendas.append([1001 + i, inicio + timedelta(days=random.randint(0, 91)), vend, reg, prod, qtd, preco, qtd * preco])
vendas.sort(key=lambda v: v[1])
for i, v in enumerate(vendas):
    v[0] = 1001 + i
ws = aba(wb, "Vendas", ["Pedido", "Data", "Vendedor", "Região", "Produto", "Quantidade", "Preço Unitário", "Total"], vendas)
for r in range(2, 42):
    ws[f"B{r}"].number_format = "DD/MM/YYYY"
    ws[f"G{r}"].number_format = ws[f"H{r}"].number_format = '"R$" #,##0.00'

# Pedidos
pedidos = [[5001 + i, random.choice(produtos)[0], random.randint(1, 10), None, None, None] for i in range(20)]
ws = aba(wb, "Pedidos", ["Pedido", "Código", "Quantidade", "Produto", "Preço", "Total"], pedidos, colunas_fazer="DEF")

# Limpeza
limpeza = [
    ["  maria DA silva ", "maria.silva@gmail.com", 12345678909],
    ["JOÃO pereira", "joao.p@empresa.com.br", 98765432100],
    [" carlos   eduardo santos", "carlos@outlook.com", 1234567890],
    ["ana PAULA  costa  ", "anapaula@yahoo.com.br", 45678912345],
    ["pedro henrique", "pedro.h@escola.edu.br", 32165498700],
    ["  LUCIANA ferreira", "lu.ferreira@gmail.com", 74185296300],
    ["rafael   alves", "rafael@empresa.com.br", 15975345682],
    ["BEATRIZ gomes ", "bia.gomes@hotmail.com", 85296374100],
]
ws = aba(wb, "Limpeza", ["Nome (bagunçado)", "E-mail", "CPF (só números)", "Nome Limpo", "Primeiro Nome", "Domínio", "CPF Formatado"],
         [l + [None] * 4 for l in limpeza], colunas_fazer="DEFG")
ws.column_dimensions["A"].width = ws.column_dimensions["B"].width = 28

# Datas
nomes = ["Ana Souza", "Bruno Lima", "Carla Dias", "Diego Alves", "Elisa Rocha", "Fábio Nunes", "Gabriela Melo", "Heitor Costa"]
datas = [[n, date(random.randint(2012, 2026), random.randint(1, 12), random.randint(1, 28)),
          date(random.randint(1970, 2004), random.randint(1, 12), random.randint(1, 28)), None, None, None] for n in nomes]
for d in datas:
    d[1] = min(d[1], date(2026, 9, 1))
ws = aba(wb, "Datas", ["Funcionário", "Admissão", "Nascimento", "Idade", "Anos de Empresa", "Fim da Experiência"],
         datas, colunas_fazer="DEF")
for r in range(2, 10):
    ws[f"B{r}"].number_format = ws[f"C{r}"].number_format = ws[f"F{r}"].number_format = "DD/MM/YYYY"

SAIDA.parent.mkdir(exist_ok=True)
wb.save(SAIDA)
print("Gerado:", SAIDA)
